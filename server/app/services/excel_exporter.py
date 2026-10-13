"""
Экспорт данных: CSV логи, Excel аналитика/комментарии.
Зависит от: database, models, openpyxl
"""

import csv
import io
from datetime import datetime, timedelta
from typing import Optional
from uuid import UUID

import structlog
from openpyxl import Workbook
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from app.database import async_session_factory
from app.models.log import ActionLog
from app.models.comment import TrackedComment
from app.models.post import Post
from app.models.account import SocialAccount

logger = structlog.get_logger()


class ExcelExporter:

    async def export_logs_csv(self, days: int = 7, limit: int = 10000) -> str:
        """Экспорт логов в CSV"""
        since = datetime.utcnow() - timedelta(days=days)
        async with async_session_factory() as db:
            result = await db.execute(
                select(ActionLog)
                .where(ActionLog.created_at >= since)
                .order_by(ActionLog.created_at.desc())
                .limit(limit)
            )
            logs = result.scalars().all()

        output = io.StringIO()
        writer = csv.writer(output)
        writer.writerow(["ID", "Device ID", "Account ID", "Action Type", "Success", "Duration (ms)", "Created At", "Details"])
        for log in logs:
            writer.writerow([
                log.id,
                str(log.device_id or ""),
                str(log.account_id or ""),
                log.action_type,
                log.success,
                log.duration_ms,
                log.created_at.isoformat() if log.created_at else "",
                str(log.details),
            ])

        logger.info("export.logs_csv", count=len(logs), days=days)
        return output.getvalue()

    async def export_comments_excel(self, user_id: Optional[UUID] = None) -> bytes:
        """Экспорт комментариев в Excel"""
        async with async_session_factory() as db:
            query = select(TrackedComment).order_by(TrackedComment.created_at.desc()).limit(5000)
            result = await db.execute(query)
            comments = result.scalars().all()

        wb = Workbook()
        ws = wb.active
        ws.title = "Комментарии"

        headers = [
            "Platform", "Post URL", "Commenter", "Gender", "Country",
            "Comment", "Response Status", "Response Text", "Response Style",
            "Created At",
        ]
        ws.append(headers)

        for col in range(1, len(headers) + 1):
            ws.cell(row=1, column=col).font = ws.cell(row=1, column=col).font.copy(bold=True)

        for c in comments:
            ws.append([
                c.platform,
                c.post_url or "",
                c.commenter_username or "",
                c.commenter_gender or "",
                c.commenter_country or "",
                c.comment_text or "",
                c.response_status,
                c.response_text or "",
                c.response_style or "",
                c.created_at.isoformat() if c.created_at else "",
            ])

        for col in ws.columns:
            max_len = max(len(str(cell.value or "")) for cell in col)
            ws.column_dimensions[col[0].column_letter].width = min(max_len + 2, 50)

        output = io.BytesIO()
        wb.save(output)
        logger.info("export.comments_excel", count=len(comments))
        return output.getvalue()

    async def export_analytics_excel(self, user_id: UUID, days: int = 30) -> bytes:
        """Экспорт аналитики в Excel"""
        since = datetime.utcnow() - timedelta(days=days)

        async with async_session_factory() as db:
            posts_result = await db.execute(
                select(Post)
                .where(Post.user_id == user_id, Post.created_at >= since)
                .order_by(Post.created_at.desc())
            )
            posts = posts_result.scalars().all()

            accounts_result = await db.execute(
                select(SocialAccount).where(SocialAccount.user_id == user_id)
            )
            accounts = accounts_result.scalars().all()

        wb = Workbook()

        # Sheet 1: Posts
        ws_posts = wb.active
        ws_posts.title = "Посты"
        ws_posts.append(["ID", "Platform", "Account", "Caption", "Status", "Photos Count", "Scheduled At", "Published At", "Created At"])
        for p in posts:
            ws_posts.append([
                str(p.id), p.platform or "",
                str(p.account_id or ""),
                (p.caption or "")[:100],
                p.status,
                len(p.photos) if p.photos else 0,
                p.scheduled_at.isoformat() if p.scheduled_at else "",
                p.published_at.isoformat() if p.published_at else "",
                p.created_at.isoformat() if p.created_at else "",
            ])

        # Sheet 2: Accounts
        ws_accounts = wb.create_sheet("Аккаунты")
        ws_accounts.append(["Platform", "Username", "Status", "Followers", "Following", "Posts", "Warmup", "Login Status"])
        for a in accounts:
            ws_accounts.append([
                a.platform, a.username, a.status,
                a.followers_count, a.following_count, a.posts_count,
                "Да" if a.warmup_enabled else "Нет",
                a.login_status,
            ])

        output = io.BytesIO()
        wb.save(output)
        logger.info("export.analytics_excel", posts=len(posts), accounts=len(accounts))
        return output.getvalue()


excel_exporter = ExcelExporter()
