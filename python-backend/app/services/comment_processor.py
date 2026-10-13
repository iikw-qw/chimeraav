"""Process incoming comment events — analyze, persist, export, broadcast."""

from __future__ import annotations

import asyncio
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook

from app.api._hub import hub
from app.config import settings
from app.database import SessionLocal
from app.models.comment import Comment
from app.services.account_analyzer import account_analyzer

_EXCEL_HEADERS = (
    "id",
    "created_at",
    "account_id",
    "platform",
    "post_id",
    "username",
    "gender",
    "language",
    "quality",
    "followers",
    "following",
    "posts",
    "account_age_days",
    "ratio",
    "has_avatar",
    "verified",
    "comment_text",
)


class CommentProcessor:
    """Pipeline: analyze → persist (if should_save) → export → broadcast."""

    def __init__(self) -> None:
        self._excel_lock = asyncio.Lock()

    async def process(self, event: dict[str, Any]) -> dict[str, Any]:
        profile = event.get("profile") or event.get("profile_data") or {}
        comment_text = event.get("comment_text") or event.get("text")

        analysis = account_analyzer.analyze(profile, comment_text)

        if not analysis["should_save"]:
            return {"saved": False, "analysis": analysis}

        comment = await self._save_comment(event, analysis)
        await self._append_to_excel(comment, analysis)
        await hub.broadcast("comments", self._payload(comment, analysis))

        return {"saved": True, "comment_id": comment.id, "analysis": analysis}

    async def _save_comment(
        self, event: dict[str, Any], analysis: dict[str, Any]
    ) -> Comment:
        async with SessionLocal() as session:
            comment = Comment(
                account_id=event["account_id"],
                platform=event.get("platform", "instagram"),
                post_id_external=event.get("post_id") or event.get("post_id_external"),
                commenter_username=event.get("commenter_username")
                or event.get("username"),
                commenter_user_id=event.get("commenter_user_id")
                or event.get("user_id"),
                commenter_gender=analysis["gender"],
                commenter_language=analysis["language"],
                commenter_metrics=analysis["metrics"],
                commenter_profile_analysis={
                    "quality": analysis["quality"],
                    "reason": analysis["reason"],
                },
                comment_text=event.get("comment_text") or event.get("text"),
            )
            session.add(comment)
            await session.commit()
            await session.refresh(comment)
            return comment

    async def _append_to_excel(
        self, comment: Comment, analysis: dict[str, Any]
    ) -> None:
        exports_dir = Path(settings.data_dir) / "exports"
        exports_dir.mkdir(parents=True, exist_ok=True)
        day = datetime.now(timezone.utc).strftime("%Y%m%d")
        path = exports_dir / f"comments_{day}.xlsx"

        metrics = analysis.get("metrics") or {}
        row = [
            comment.id,
            comment.created_at.isoformat() if comment.created_at else None,
            comment.account_id,
            comment.platform,
            comment.post_id_external,
            comment.commenter_username,
            analysis.get("gender"),
            analysis.get("language"),
            analysis.get("quality"),
            metrics.get("followers"),
            metrics.get("following"),
            metrics.get("posts"),
            metrics.get("account_age_days"),
            metrics.get("ratio"),
            metrics.get("has_avatar"),
            metrics.get("verified"),
            comment.comment_text,
        ]

        async with self._excel_lock:
            await asyncio.to_thread(self._write_row, path, row)

    @staticmethod
    def _write_row(path: Path, row: list[Any]) -> None:
        if path.exists():
            wb = load_workbook(path)
            ws = wb.active
        else:
            wb = Workbook()
            ws = wb.active
            ws.title = "comments"
            ws.append(list(_EXCEL_HEADERS))
        ws.append(row)
        wb.save(path)

    @staticmethod
    def _payload(comment: Comment, analysis: dict[str, Any]) -> dict[str, Any]:
        return {
            "id": comment.id,
            "account_id": comment.account_id,
            "platform": comment.platform,
            "username": comment.commenter_username,
            "gender": analysis.get("gender"),
            "language": analysis.get("language"),
            "quality": analysis.get("quality"),
            "metrics": analysis.get("metrics"),
            "comment_text": comment.comment_text,
            "created_at": comment.created_at.isoformat() if comment.created_at else None,
        }


comment_processor = CommentProcessor()
